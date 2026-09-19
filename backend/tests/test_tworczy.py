"""Testy modułów twórczych: strony WWW, obrazy (tło, gumka), tłumaczenie dokumentów."""

from __future__ import annotations

import io
import json
import time
from collections.abc import Iterator
from pathlib import Path

import numpy as np
import pymupdf
import pytest
from conftest import FONT_FILE, ToolHarness
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from PIL import Image
from test_api import HEADERS, login, set_password

from nexus.api.app import create_app
from nexus.api.modules import strony as strony_api
from nexus.api.modules import tlumacz as tlumacz_api
from nexus.config import Settings
from nexus.tools import registry
from nexus.tools import tlumacz as tlumacz_tool
from nexus.tools.base import ToolError, ToolResult
from nexus.tworczy import obrazy, tlumaczenie
from nexus.tworczy.strony import SiteError, SiteStore, check_address, check_path


def call(harness: ToolHarness, tool_name: str, /, **arguments: object) -> ToolResult:
    tool = registry.get(tool_name)
    return tool.handler(harness.context(), tool.parse(arguments))


def fake_translate(texts: list[str]) -> list[str]:
    """Atrapa tłumacza: wielkie litery z zachowaniem znaczników fragmentów."""
    return [text.upper() for text in texts]


# --- strony: bezpieczne ścieżki i magazyn ---------------------------------------------------------


@pytest.mark.parametrize(
    "path",
    [
        "../etc/passwd",
        "a/../../b.html",
        "/etc/passwd",
        "C:/x.html",
        "a\\b.html",
        ".env",
        "css/.ukryty.css",
        "a\x00.html",
    ],
)
def test_site_paths_reject_escapes(path: str) -> None:
    with pytest.raises(SiteError):
        check_path(path)


def test_site_paths_and_addresses_are_normalized() -> None:
    assert check_path("./css//style.css") == "css/style.css"
    assert check_address(" Kawiarnia-2026 ") == "kawiarnia-2026"
    for bad in ("-zle", "zle-", "a" * 49, "ze spacją", ".meta", "..", "strona_1"):
        with pytest.raises(SiteError):
            check_address(bad)


def test_site_store_files_versions_and_publication(tmp_path: Path) -> None:
    store = SiteStore(tmp_path / "strony", max_site_mb=1)
    store.create("kawiarnia", "Kawiarnia <Pod Lipą>")
    assert "&lt;Pod Lipą&gt;" in store.read_text("kawiarnia", "index.html")
    with pytest.raises(SiteError, match="już istnieje"):
        store.create("kawiarnia", "Druga")
    store.write_text("kawiarnia", "index.html", "<h1>Wersja 1</h1>")
    store.write_text("kawiarnia", "css/style.css", "h1{color:red}")
    with pytest.raises(SiteError, match="Niedozwolony typ"):
        store.write_text("kawiarnia", "skrypt.php", "<?php ?>")
    with pytest.raises(SiteError, match="limit"):
        store.write_bytes("kawiarnia", "img/duzy.png", b"0" * (2 * 1024 * 1024))
    assert [item["path"] for item in store.list_files("kawiarnia")] == ["css/style.css", "index.html"]

    version = store.snapshot("kawiarnia", "pierwsza")
    store.write_text("kawiarnia", "index.html", "<h1>Wersja 2</h1>")
    assert store.resolve("kawiarnia", "", published=False).read_text(encoding="utf-8") == "<h1>Wersja 2</h1>"
    assert store.resolve("kawiarnia", "", published=True) is None

    store.publish("kawiarnia")
    store.write_text("kawiarnia", "index.html", "<h1>Szkic 3</h1>")
    published = store.resolve("kawiarnia", "index.html", published=True)
    assert published.read_text(encoding="utf-8") == "<h1>Wersja 2</h1>"
    assert store.resolve("kawiarnia", "../.meta/kawiarnia.json", published=True) is None
    assert store.resolve("kawiarnia", "css", published=True) is None

    store.restore("kawiarnia", version["id"])
    assert store.read_text("kawiarnia", "index.html") == "<h1>Wersja 1</h1>"
    notes = [item["note"] for item in store.meta("kawiarnia")["versions"]]
    assert notes == ["pierwsza", "Publikacja", f"Przed przywróceniem {version['id']}"]

    store.delete_file("kawiarnia", "css/style.css")
    assert not (store.draft_dir("kawiarnia") / "css").exists()
    store.unpublish("kawiarnia")
    assert store.resolve("kawiarnia", "", published=True) is None
    store.delete_site("kawiarnia")
    assert store.list_sites() == []


def test_site_tools_write_read_and_request_publication(harness: ToolHarness, tmp_path: Path) -> None:
    SiteStore(harness.settings.data_dir / "strony").create("sklep", "Sklep")
    written = call(harness, "site_write_file", site="sklep", path="index.html", content="<p>Cześć</p>")
    assert written.data == {"path": "index.html", "size": len("<p>Cześć</p>".encode())}
    assert call(harness, "site_read_file", site="sklep", path="index.html").data["content"] == "<p>Cześć</p>"
    logo = tmp_path / "logo.png"
    Image.new("RGB", (4, 4), "red").save(logo)
    call(harness, "site_import_file", site="sklep", file_id=harness.add(logo), path="img/logo.png")
    listing = call(harness, "site_list", site="sklep").data
    assert [item["path"] for item in listing["files"]] == ["img/logo.png", "index.html"]
    with pytest.raises(ToolError, match="Nieprawidłowa ścieżka"):
        call(harness, "site_write_file", site="sklep", path="../x.html", content="x")
    with pytest.raises(ToolError, match="nie istnieje"):
        call(harness, "site_write_file", site="brak", path="index.html", content="x")

    result = call(harness, "site_publish", site="sklep", note="gotowe")
    assert result.data["status"] == "czeka_na_potwierdzenie"
    store = SiteStore(harness.settings.data_dir / "strony")
    assert store.meta("sklep")["publish_request"]["note"] == "gotowe"
    assert not store.published_dir("sklep").exists()


# --- strony: API, podgląd i publikacja ------------------------------------------------------------


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    return Settings(
        data_dir=tmp_path / "data",
        static_dir=tmp_path / "static",
        database_url=f"sqlite+aiosqlite:///{(tmp_path / 'nexus.db').as_posix()}",
        cookie_secure=False,
        redis_url="",
        voice_warm_up=False,
        voice_stt_model_dir=tmp_path / "brak-modelu",
        voice_tts_dir=tmp_path / "brak-glosow",
        voice_google_key_file=tmp_path / "brak-klucza-google",
        qdrant_url="http://127.0.0.1:1",
        realesrgan_dir=tmp_path / "brak-esrgan",
        tworczy_rembg_bin=str(tmp_path / "brak-rembg"),
    )


@pytest.fixture
def client(settings: Settings) -> Iterator[TestClient]:
    set_password(settings)
    with TestClient(create_app(settings)) as test_client:
        login(test_client)
        yield test_client


def test_slugify_polish_titles() -> None:
    assert strony_api.slugify("Kawiarnia Pod Lipą – Łódź!") == "kawiarnia-pod-lipa-lodz"
    assert strony_api.slugify("!!!") == "strona"


def test_site_api_preview_and_publication(client: TestClient, settings: Settings) -> None:
    created = client.post("/api/strony", json={"title": "Pensjonat Żuraw"}, headers=HEADERS)
    assert created.status_code == 201, created.text
    site = created.json()
    assert site["address"] == "pensjonat-zuraw"
    conversation = client.get(f"/api/conversations/{site['conversation_id']}")
    assert conversation.status_code == 200
    again = client.post("/api/strony", json={"title": "Pensjonat Żuraw"}, headers=HEADERS).json()
    assert again["address"] == "pensjonat-zuraw-2"

    store = SiteStore(settings.data_dir / "strony")
    store.write_text("pensjonat-zuraw", "index.html", "<h1>Szkic</h1>")
    store.write_text("pensjonat-zuraw", "css/style.css", "h1{}")
    detail = client.get("/api/strony/pensjonat-zuraw").json()
    assert [item["path"] for item in detail["files"]] == ["css/style.css", "index.html"]

    anonymous = TestClient(client.app)
    preview = anonymous.get(detail["preview_url"])
    assert preview.status_code == 200 and preview.text == "<h1>Szkic</h1>"
    csp = preview.headers["content-security-policy"]
    assert csp.startswith("sandbox allow-scripts") and "allow-same-origin" not in csp
    assert preview.headers["x-frame-options"] == "SAMEORIGIN"
    assert (
        anonymous.get(detail["preview_url"] + "css/style.css").headers["content-type"].startswith("text/css")
    )
    forged = detail["preview_url"].replace("/podglad/", "/podglad/9")
    assert anonymous.get(forged).status_code == 403
    other = anonymous.get(detail["preview_url"].replace("pensjonat-zuraw/", "pensjonat-zuraw-2/"))
    assert other.status_code == 403
    assert anonymous.get("/api/strony").status_code == 401

    assert anonymous.get("/s/pensjonat-zuraw/").status_code == 404
    assert client.post("/api/strony/pensjonat-zuraw/publikuj").status_code == 403
    published = client.post("/api/strony/pensjonat-zuraw/publikuj", headers=HEADERS).json()
    assert published["published_at"] and published["versions"][0]["note"] == "Publikacja"
    store.write_text("pensjonat-zuraw", "index.html", "<h1>Nowy szkic</h1>")
    public = anonymous.get("/s/pensjonat-zuraw/")
    assert public.status_code == 200 and public.text == "<h1>Szkic</h1>"
    assert public.headers["content-security-policy"] == strony_api.SITE_CSP
    assert anonymous.get("/s/pensjonat-zuraw", follow_redirects=False).status_code == 308
    assert anonymous.get("/s/pensjonat-zuraw/..%2F..%2F.klucz").status_code == 404
    assert anonymous.get("/s/pensjonat-zuraw/nie-ma.html").status_code == 404

    source = client.get("/api/strony/pensjonat-zuraw/pliki/index.html").json()
    assert source["content"] == "<h1>Nowy szkic</h1>"
    restored = client.post(
        f"/api/strony/pensjonat-zuraw/wersje/{published['versions'][0]['id']}/przywroc", headers=HEADERS
    )
    assert restored.status_code == 200
    assert store.read_text("pensjonat-zuraw", "index.html") == "<h1>Szkic</h1>"
    client.post("/api/strony/pensjonat-zuraw/wycofaj", headers=HEADERS)
    assert anonymous.get("/s/pensjonat-zuraw/").status_code == 404
    assert client.delete("/api/strony/pensjonat-zuraw", headers=HEADERS).json() == {"ok": True}
    assert client.get("/api/strony/pensjonat-zuraw").status_code == 404


def test_preview_token_expires(tmp_path: Path) -> None:
    store = SiteStore(tmp_path)
    token = strony_api.preview_token(store, "strona", now=1000)
    assert strony_api.check_preview_token(store, "strona", token, now=2000)
    assert not strony_api.check_preview_token(store, "inna", token, now=2000)
    assert not strony_api.check_preview_token(store, "strona", token, now=1000 + 13 * 3600)


# --- obrazy ---------------------------------------------------------------------------------------


def _cutout(size: tuple[int, int] = (80, 60)) -> Image.Image:
    """Czerwone koło na przezroczystym tle."""
    width, height = size
    ys, xs = np.mgrid[0:height, 0:width]
    inside = (xs - width / 2) ** 2 + (ys - height / 2) ** 2 < (min(size) / 3) ** 2
    pixels = np.zeros((height, width, 4), np.uint8)
    pixels[inside] = (220, 20, 20, 255)
    return Image.fromarray(pixels, "RGBA")


def test_compose_background_color_gradient_image_and_blur() -> None:
    cutout = _cutout()
    colored = obrazy.compose_background(cutout, "color", color="#00ff00")
    assert colored.getpixel((1, 1)) == (0, 255, 0, 255)
    assert colored.getpixel((40, 30)) == (220, 20, 20, 255)

    gradient = obrazy.compose_background(cutout, "gradient", color="#000000", color2="#ffffff", angle=90)
    assert gradient.getpixel((1, 0))[0] < 20 and gradient.getpixel((1, 59))[0] > 235

    photo = Image.new("RGB", (400, 100), (0, 0, 255))
    placed = obrazy.compose_background(cutout, "image", background=photo)
    assert placed.size == cutout.size and placed.getpixel((0, 0)) == (0, 0, 255, 255)

    original = Image.new("RGB", cutout.size, (255, 255, 255))
    original.paste((0, 0, 0), (0, 0, 40, 60))
    blurred = obrazy.compose_background(cutout, "blur", original=original, blur_radius=10)
    edge = blurred.getpixel((38, 2))[0]
    assert 20 < edge < 235
    assert blurred.getpixel((40, 30)) == (220, 20, 20, 255)

    with pytest.raises(obrazy.ImageOpError):
        obrazy.compose_background(cutout, "color", color="zielony")
    assert obrazy.parse_color("#abc") == (170, 187, 204, 255)


def test_erase_with_mask_restores_surroundings() -> None:
    image = Image.new("RGB", (120, 90), (40, 160, 90))
    image.paste((250, 250, 250), (50, 35, 70, 55))
    mask = obrazy.regions_mask(image.size, [(50 / 120, 35 / 90, 20 / 120, 20 / 90)])
    result = np.asarray(obrazy.erase_with_mask(image, mask)).astype(int)
    patch = result[38:52, 53:67]
    assert np.abs(patch - np.array([40, 160, 90])).max() < 25
    with pytest.raises(obrazy.ImageOpError, match="pusta"):
        obrazy.erase_with_mask(image, Image.new("L", image.size, 0))


def test_image_tools_with_transparent_cutout(harness: ToolHarness, tmp_path: Path) -> None:
    cutout_path = tmp_path / "produkt.png"
    _cutout().save(cutout_path)
    result = call(
        harness, "change_background", file_id=harness.add(cutout_path), mode="color", color="#ffffff"
    )
    output = Image.open(result.files[0].path)
    assert result.files[0].name == "produkt_nowe_tlo.png"
    assert output.convert("RGB").getpixel((0, 0)) == (255, 255, 255)

    photo = tmp_path / "zdjecie.jpg"
    image = Image.new("RGB", (100, 80), (30, 60, 200))
    image.paste((255, 255, 0), (10, 10, 30, 20))
    image.save(photo, quality=95)
    erased = call(harness, "erase_objects", file_id=harness.add(photo), regions=[[0.1, 0.125, 0.2, 0.125]])
    assert erased.files[0].name == "zdjecie_gumka.jpg"
    repaired = np.asarray(Image.open(erased.files[0].path)).astype(int)
    assert np.abs(repaired[14, 20] - np.array([30, 60, 200])).max() < 40

    opaque = tmp_path / "pelne.png"
    Image.new("RGB", (10, 10)).save(opaque)
    with pytest.raises(ToolError, match="rembg"):
        harness.settings.tworczy_rembg_bin = str(tmp_path / "brak")
        call(harness, "remove_background", file_ids=[harness.add(opaque)])


def _wait(client: TestClient, url: str) -> dict:
    for _ in range(100):
        state = client.get(url).json()
        if state["status"] != "running":
            return state
        time.sleep(0.05)
    raise AssertionError("Zadanie nie zakończyło się.")


def test_image_api_jobs(client: TestClient, tmp_path: Path) -> None:
    buffer = io.BytesIO()
    _cutout().save(buffer, format="PNG")
    uploaded = client.post("/api/files", files={"file": ("wycinek.png", buffer.getvalue())}, headers=HEADERS)
    file_id = uploaded.json()["id"]
    assert client.get("/api/obrazy/mozliwosci").json() == {"remove_background": False, "upscale": False}

    job = client.post(
        "/api/obrazy/zmien-tlo",
        json={"file_id": file_id, "mode": "gradient", "color": "#101010", "color2": "#fafafa"},
        headers=HEADERS,
    )
    assert job.status_code == 202, job.text
    state = _wait(client, f"/api/obrazy/zadania/{job.json()['id']}")
    assert state["status"] == "done", state
    result_file = state["result"]["files"][0]
    assert result_file["name"] == "wycinek_nowe_tlo.png"
    assert client.get(f"/api/files/{result_file['id']}/download").status_code == 200

    mask = Image.new("RGBA", (80, 60), (0, 0, 0, 0))
    mask.paste((255, 0, 0, 255), (35, 25, 45, 35))
    mask_buffer = io.BytesIO()
    mask.save(mask_buffer, format="PNG")
    import base64

    data_url = "data:image/png;base64," + base64.b64encode(mask_buffer.getvalue()).decode()
    erase = client.post("/api/obrazy/gumka", json={"file_id": file_id, "mask": data_url}, headers=HEADERS)
    assert _wait(client, f"/api/obrazy/zadania/{erase.json()['id']}")["status"] == "done"
    assert not list((client.app.state.settings.work_dir / "tworczy-maski").glob("*.png"))

    bad = client.post("/api/obrazy/gumka", json={"file_id": file_id, "mask": "nie-png"}, headers=HEADERS)
    assert bad.status_code == 422
    missing = client.post(
        "/api/obrazy/usun-tlo", json={"file_id": "00000000-0000-0000-0000-000000000000"}, headers=HEADERS
    )
    assert missing.status_code == 404
    assert client.get("/api/obrazy/zadania/nieznane").status_code == 404


# --- tłumaczenie ----------------------------------------------------------------------------------


def test_parts_encoding_and_fallback() -> None:
    assert tlumaczenie.encode_parts(["Hello ", "world"]) == "<1>Hello </1><2>world</2>"
    assert tlumaczenie.decode_parts("<2>świecie</2><1>Witaj </1>", 2) == ["Witaj ", "świecie"]
    assert tlumaczenie.decode_parts("<1>Witaj</1> świecie", 2) is None
    assert tlumaczenie.distribute("<1>Witaj</1> świecie", 2) == ["Witaj świecie", ""]
    assert tlumaczenie.batches(["a" * 10, "b" * 10, "c" * 10], 25) == [[0, 1], [2]]


def test_parse_cli_output_variants() -> None:
    structured = json.dumps(
        {"type": "result", "structured_output": {"translations": ["Hallo"], "source_language": "pl"}}
    )
    assert tlumaczenie.parse_cli_output(structured, 1) == (["Hallo"], "pl")
    fenced = json.dumps({"result": '```json\n{"translations": ["A", "B"]}\n```'})
    assert tlumaczenie.parse_cli_output(fenced, 2) == (["A", "B"], "")
    with pytest.raises(tlumaczenie.TranslationError, match="zamiast"):
        tlumaczenie.parse_cli_output(fenced, 3)
    with pytest.raises(tlumaczenie.TranslationError, match="błąd"):
        tlumaczenie.parse_cli_output(json.dumps({"is_error": True, "result": "limit"}), 1)


def test_translate_segments_deduplicates_and_retries_single() -> None:
    calls: list[list[str]] = []

    def flaky(texts: list[str]) -> list[str]:
        calls.append(texts)
        if len(texts) > 1:
            raise tlumaczenie.TranslationError("zła liczba segmentów")
        return [f"[{texts[0]}]"]

    result = tlumaczenie.translate_segments(["jeden", "dwa", "jeden", "123", "  "], flaky)
    assert result == ["[jeden]", "[dwa]", "[jeden]", "123", "  "]
    assert calls == [["jeden", "dwa"], ["jeden"], ["dwa"]]


def test_claude_translator_command_and_batches(settings: Settings) -> None:
    commands: list[list[str]] = []

    def runner(command: list[str], env: dict[str, str], cwd: Path) -> str:
        commands.append(command)
        segments = json.loads(command[2].split("Segmenty (JSON):\n", 1)[1])["segments"]
        answer = {"translations": [text.upper() for text in segments], "source_language": "pl"}
        return json.dumps({"type": "result", "structured_output": answer})

    settings.tworczy_translate_batch_chars = 12
    translator = tlumaczenie.ClaudeTranslator(settings, "en", runner=runner)
    assert translator(["Dzień dobry", "Do widzenia", "Dzień dobry"]) == [
        "DZIEŃ DOBRY",
        "DO WIDZENIA",
        "DZIEŃ DOBRY",
    ]
    assert len(commands) == 2 and translator.detected == "pl"
    command = commands[0]
    assert command[command.index("--tools") + 1] == ""
    assert (
        "--no-session-persistence" in command and command[command.index("--model") + 1] == "claude-sonnet-5"
    )
    with pytest.raises(tlumaczenie.TranslationError, match="Nieobsługiwany"):
        tlumaczenie.ClaudeTranslator(settings, "klingoński")


def test_translate_docx_keeps_run_formatting(tmp_path: Path) -> None:
    source = tmp_path / "umowa.docx"
    document = DocxDocument()
    document.sections[0].header.paragraphs[0].text = "Nagłówek firmy"
    paragraph = document.add_paragraph("Strony ustalają ")
    paragraph.add_run("cenę").bold = True
    paragraph.add_run(" usługi.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Kwota"
    table.cell(0, 1).text = "100 zł"
    document.add_paragraph("2026")
    document.save(source)

    target = tmp_path / "umowa_en.docx"
    count = tlumaczenie.translate_docx(source, target, fake_translate)
    assert count == 4
    result = DocxDocument(str(target))
    runs = result.paragraphs[0].runs
    assert [run.text for run in runs] == ["STRONY USTALAJĄ ", "CENĘ", " USŁUGI."]
    assert runs[1].bold is True and not runs[0].bold
    assert result.tables[0].cell(0, 0).text == "KWOTA"
    assert result.tables[0].cell(0, 1).text == "100 ZŁ"
    assert result.sections[0].header.paragraphs[0].text == "NAGŁÓWEK FIRMY"
    assert result.paragraphs[1].text == "2026"


def test_translate_pptx_and_plain_text(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    source = tmp_path / "oferta.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Oferta pokoi"
    body = slide.placeholders[1].text_frame
    body.text = "Śniadanie w cenie"
    run = body.paragraphs[0].add_run()
    run.text = " – promocja"
    run.font.bold = True
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(4), Inches(3), Inches(1)).table
    table.cell(0, 0).text = "Pokój dwuosobowy"
    slide.notes_slide.notes_text_frame.text = "Notatka prelegenta"
    presentation.save(str(source))

    target = tmp_path / "oferta_en.pptx"
    assert tlumaczenie.translate_pptx(source, target, fake_translate) == 4
    result = Presentation(str(target)).slides[0]
    assert result.shapes.title.text == "OFERTA POKOI"
    runs = result.placeholders[1].text_frame.paragraphs[0].runs
    assert [item.text for item in runs] == ["ŚNIADANIE W CENIE", " – PROMOCJA"] and runs[1].font.bold
    assert "POKÓJ DWUOSOBOWY" in [shape.table.cell(0, 0).text for shape in result.shapes if shape.has_table]
    assert result.notes_slide.notes_text_frame.text == "NOTATKA PRELEGENTA"

    text = tmp_path / "notatka.md"
    text.write_text("# Tytuł\n\nPierwszy akapit.\n\n123\n", encoding="utf-8")
    tlumaczenie.translate_plain(text, tmp_path / "notatka_en.md", fake_translate)
    assert (tmp_path / "notatka_en.md").read_text(encoding="utf-8") == "# TYTUŁ\n\nPIERWSZY AKAPIT.\n\n123\n"


def test_translate_pdf_replaces_text_in_place(tmp_path: Path) -> None:
    source = tmp_path / "pismo.pdf"
    with pymupdf.open() as document:
        page = document.new_page(width=595, height=842)
        page.draw_rect(pymupdf.Rect(40, 40, 555, 200), color=(0, 0, 1), fill=(0.9, 0.9, 1))
        page.insert_font(fontname="doc", fontfile=FONT_FILE)
        page.insert_textbox(pymupdf.Rect(60, 60, 540, 180), "Zażółć gęślą jaźń", fontname="doc", fontsize=14)
        page.insert_textbox(pymupdf.Rect(60, 400, 540, 500), "Drugi akapit", fontname="doc", fontsize=11)
        document.save(source)
    target = tmp_path / "pismo_en.pdf"
    assert tlumaczenie.translate_pdf(source, target, fake_translate) == 2
    with pymupdf.open(target) as document:
        text = document[0].get_text().replace("\xa0", " ")
        blocks = tlumaczenie.pdf_blocks(document[0])
        assert len(document[0].get_drawings()) == 1
    assert "ZAŻÓŁĆ GĘŚLĄ JAŹŃ" in text and "Zażółć" not in text and "DRUGI AKAPIT" in text
    first = next(block for block in blocks if "ZAŻÓŁĆ" in block["text"])
    assert 55 <= first["bbox"][0] <= 65 and 55 <= first["bbox"][1] <= 70
    assert first["size"] == pytest.approx(14, abs=1.5)

    scanned = tmp_path / "skan.pdf"
    with pymupdf.open() as document:
        document.new_page()
        document.save(scanned)
    with pytest.raises(tlumaczenie.TranslationError, match="OCR"):
        tlumaczenie.translate_pdf(scanned, tmp_path / "x.pdf", fake_translate)


def test_translate_document_tool_with_fake_translator(
    harness: ToolHarness, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "list.docx"
    document = DocxDocument()
    document.add_paragraph("Dziękujemy za rezerwację.")
    document.save(source)
    monkeypatch.setattr(tlumacz_tool, "tlumaczenie_factory", lambda ctx, args: fake_translate)
    result = call(harness, "translate_document", file_id=harness.add(source), target_language="en")
    assert result.files[0].name == "list_en.docx"
    assert DocxDocument(str(result.files[0].path)).paragraphs[0].text == "DZIĘKUJEMY ZA REZERWACJĘ."
    with pytest.raises(ToolError, match="Nieobsługiwany język"):
        call(harness, "translate_document", file_id=harness.add(source), target_language="xx")
    odt = tmp_path / "list.odt"
    odt.write_bytes(b"x")
    with pytest.raises(ToolError, match="Nieobsługiwany format"):
        call(harness, "translate_document", file_id=harness.add(odt), target_language="en")


def test_translation_api(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    class Fake:
        detected = "pl"

        def __call__(self, texts: list[str]) -> list[str]:
            return fake_translate(texts)

    monkeypatch.setattr(tlumacz_api, "translator_for", lambda request, payload: Fake())
    languages = client.get("/api/tlumacz/jezyki").json()
    assert {"code": "en", "name": "angielski"} in languages["languages"]
    response = client.post(
        "/api/tlumacz/tekst", json={"text": "Dzień dobry", "target": "en"}, headers=HEADERS
    )
    assert response.json() == {"translation": "DZIEŃ DOBRY", "source_language": "pl"}
    assert (
        client.post("/api/tlumacz/tekst", json={"text": "", "target": "en"}, headers=HEADERS).status_code
        == 422
    )

    monkeypatch.setattr(tlumacz_tool, "tlumaczenie_factory", lambda ctx, args: fake_translate)
    uploaded = client.post(
        "/api/files", files={"file": ("notatka.txt", "Spotkanie o dziesiątej.".encode())}, headers=HEADERS
    ).json()
    job = client.post(
        "/api/tlumacz/dokument", json={"file_id": uploaded["id"], "target": "de"}, headers=HEADERS
    )
    state = _wait(client, f"/api/tlumacz/zadania/{job.json()['id']}")
    assert state["status"] == "done", state
    produced = state["result"]["files"][0]
    assert produced["name"] == "notatka_de.txt"
    assert client.get(f"/api/files/{produced['id']}/download").text == "SPOTKANIE O DZIESIĄTEJ."
