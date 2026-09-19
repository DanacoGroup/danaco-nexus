"""Narzędzia dokumentów biurowych: konwersje (LibreOffice, Inkscape), tworzenie
dokumentów z treści przygotowanej przez agenta, korekta językowa (LanguageTool)."""

from __future__ import annotations

import csv
import io
import re
from typing import Any, Literal

import httpx
from docx import Document as DocxDocument
from docx.shared import Pt
from pydantic import Field

from nexus.storage import safe_filename
from nexus.tools.base import (
    OutputFile,
    ToolContext,
    ToolError,
    ToolInput,
    ToolResult,
    registry,
    truncate_text,
)
from nexus.tools.common import file_kind, libreoffice_convert, unique_name, with_suffix
from nexus.tools.files import _tika_text

OFFICE_TARGETS = {
    "pdf": "pdf",
    "docx": "docx:MS Word 2007 XML",
    "odt": "odt",
    "rtf": "rtf",
    "txt": "txt:Text",
    "html": "html",
    "xlsx": "xlsx:Calc MS Excel 2007 XML",
    "ods": "ods",
    "csv": "csv",
    "pptx": "pptx:Impress MS PowerPoint 2007 XML",
    "odp": "odp",
}
LANGUAGETOOL_CHUNK = 18_000


class ConvertInput(ToolInput):
    file_ids: list[str] = Field(min_length=1, max_length=100, description="Pliki do konwersji.")
    target_format: Literal[
        "pdf", "docx", "odt", "rtf", "txt", "html", "xlsx", "ods", "csv", "pptx", "odp", "png"
    ] = Field(description="Format docelowy.")


@registry.register(
    "convert_documents",
    """Konwertuje dokumenty pakietem LibreOffice (DOC/DOCX/ODT/RTF/XLS/XLSX/ODS/CSV/PPT/PPTX/
ODP/HTML/TXT ⇄ PDF, DOCX, ODT, XLSX itd.) oraz grafikę wektorową SVG do PDF/PNG (Inkscape).""",
    ConvertInput,
)
def convert_documents(ctx: ToolContext, args: ConvertInput) -> ToolResult:
    used: set[str] = set()
    outputs: list[OutputFile] = []
    for file_id in args.file_ids:
        ctx.check_cancelled()
        file = ctx.file(file_id)
        kind = file_kind(file)
        target_name = unique_name(with_suffix(file.name, f".{args.target_format}"), used)
        if kind == "svg":
            if args.target_format not in {"pdf", "png"}:
                raise ToolError("SVG można skonwertować do PDF albo PNG.")
            target = ctx.output_path(target_name)
            ctx.run_command(
                [
                    "inkscape",
                    str(file.path),
                    f"--export-type={args.target_format}",
                    f"--export-filename={target}",
                ],
                timeout=300,
            )
        elif kind in {"office", "text", "pdf"}:
            if args.target_format == "png":
                raise ToolError("Dokument do PNG: użyj view_pages albo convert do PDF.")
            staged = ctx.output_path(safe_filename(file.name))
            staged.write_bytes(file.path.read_bytes())
            produced = libreoffice_convert(ctx, staged, OFFICE_TARGETS[args.target_format])
            target = produced.with_name(target_name)
            produced.rename(target)
        else:
            raise ToolError(f"{file.name}: nieobsługiwany rodzaj pliku do konwersji dokumentów.")
        outputs.append(OutputFile(target, target.name, f"Konwersja {file.name}"))
    return ToolResult(
        {"outputs": [f.name for f in outputs]}, f"Skonwertowano dokumenty: {len(outputs)}", files=outputs
    )


class WriteDocumentInput(ToolInput):
    name: str = Field(description="Nazwa pliku bez rozszerzenia, np. 'Podsumowanie umowy'.")
    format: Literal["docx", "pdf", "md", "txt", "xlsx", "csv"] = Field(description="Format pliku.")
    content: str = Field(
        description="Treść w Markdown (nagłówki #, listy -, **pogrubienie**); "
        "dla xlsx/csv: dane CSV z nagłówkiem w pierwszym wierszu."
    )


def _add_runs(paragraph: Any, text: str) -> None:
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part:
            paragraph.add_run(part)


def markdown_to_docx(content: str) -> DocxDocument:
    """Prosty konwerter Markdown → DOCX (nagłówki, listy, pogrubienia, akapity)."""
    document = DocxDocument()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)
    for raw in content.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)", line)
        bullet = re.match(r"^\s*[-*•]\s+(.*)", line)
        numbered = re.match(r"^\s*\d+[.)]\s+(.*)", line)
        if heading:
            document.add_heading(heading.group(2).strip(), level=len(heading.group(1)))
        elif bullet:
            _add_runs(document.add_paragraph(style="List Bullet"), bullet.group(1))
        elif numbered:
            _add_runs(document.add_paragraph(style="List Number"), numbered.group(1))
        else:
            _add_runs(document.add_paragraph(), line)
    return document


@registry.register(
    "write_document",
    """Tworzy plik z treści przygotowanej przez Ciebie: raport, podsumowanie, pismo,
zestawienie. DOCX/PDF z formatowaniem Markdown, MD, TXT, a tabele jako XLSX/CSV.""",
    WriteDocumentInput,
)
def write_document(ctx: ToolContext, args: WriteDocumentInput) -> ToolResult:
    name = safe_filename(args.name, "dokument")
    if args.format in {"md", "txt"}:
        target = ctx.output_path(f"{name}.{args.format}")
        target.write_text(args.content, encoding="utf-8")
    elif args.format in {"csv", "xlsx"}:
        rows = list(csv.reader(io.StringIO(args.content.strip())))
        if not rows:
            raise ToolError("Brak danych tabeli.")
        csv_path = ctx.output_path(f"{name}.csv")
        with csv_path.open("w", encoding="utf-8", newline="") as handle:
            csv.writer(handle).writerows(rows)
        target = (
            csv_path
            if args.format == "csv"
            else libreoffice_convert(ctx, csv_path, "xlsx:Calc MS Excel 2007 XML:44,34,76,1")
        )
    else:
        docx_path = ctx.output_path(f"{name}.docx")
        markdown_to_docx(args.content).save(str(docx_path))
        target = docx_path if args.format == "docx" else libreoffice_convert(ctx, docx_path, "pdf")
    return ToolResult(
        {"output": target.name},
        f"Utworzono {target.name}",
        files=[OutputFile(target, target.name, "Dokument utworzony przez asystenta")],
    )


class GrammarInput(ToolInput):
    text: str | None = Field(None, description="Tekst do sprawdzenia (albo użyj file_id).")
    file_id: str | None = Field(None, description="Plik tekstowy/DOCX/PDF z warstwą tekstową.")
    language: str = Field("pl-PL", description="Kod języka LanguageTool, np. 'pl-PL', 'en-US', 'auto'.")
    apply_corrections: bool = Field(False, description="Zwróć tekst z zastosowanymi pierwszymi sugestiami.")


def _check_chunk(ctx: ToolContext, text: str, language: str) -> list[dict[str, Any]]:
    try:
        response = httpx.post(
            f"{ctx.settings.languagetool_url}/v2/check",
            data={"text": text, "language": language},
            timeout=180,
        )
    except httpx.HTTPError as error:
        raise ToolError(f"Usługa LanguageTool jest niedostępna: {error}") from error
    if response.status_code != 200:
        raise ToolError(f"LanguageTool zwrócił błąd HTTP {response.status_code}: {response.text[:200]}")
    return response.json().get("matches", [])


@registry.register(
    "check_grammar",
    """Sprawdza pisownię, gramatykę, interpunkcję i styl (LanguageTool, domyślnie polski).
Zwraca listę uwag z propozycjami poprawek i opcjonalnie poprawiony tekst.""",
    GrammarInput,
)
def check_grammar(ctx: ToolContext, args: GrammarInput) -> ToolResult:
    if args.text:
        text = args.text
    elif args.file_id:
        file = ctx.file(args.file_id)
        text = (
            file.path.read_text(encoding="utf-8", errors="replace")
            if file_kind(file) == "text"
            else _tika_text(ctx, file)
        )
    else:
        raise ToolError("Podaj text albo file_id.")
    issues: list[dict[str, Any]] = []
    corrected_parts: list[str] = []
    offset_base = 0
    for start in range(0, len(text), LANGUAGETOOL_CHUNK):
        ctx.check_cancelled()
        chunk = text[start : start + LANGUAGETOOL_CHUNK]
        matches = _check_chunk(ctx, chunk, args.language)
        corrected = chunk
        for match in sorted(matches, key=lambda item: item["offset"], reverse=True):
            replacements = [r["value"] for r in match.get("replacements", [])[:3]]
            issues.append(
                {
                    "offset": offset_base + match["offset"],
                    "context": match.get("context", {}).get("text", "")[:120],
                    "problem": chunk[match["offset"] : match["offset"] + match["length"]],
                    "message": match.get("message", ""),
                    "suggestions": replacements,
                    "rule": match.get("rule", {}).get("category", {}).get("name", ""),
                }
            )
            if args.apply_corrections and replacements:
                corrected = (
                    corrected[: match["offset"]]
                    + replacements[0]
                    + corrected[match["offset"] + match["length"] :]
                )
        corrected_parts.append(corrected)
        offset_base += len(chunk)
    issues.sort(key=lambda item: item["offset"])
    data: dict[str, Any] = {"issues_count": len(issues), "issues": issues[:300]}
    if len(issues) > 300:
        data["note"] = "Pokazano pierwsze 300 uwag."
    if args.apply_corrections:
        data["corrected_text"], truncated = truncate_text("".join(corrected_parts))
        if truncated:
            data["corrected_text_truncated"] = True
    return ToolResult(data, f"LanguageTool: {len(issues)} uwag")
